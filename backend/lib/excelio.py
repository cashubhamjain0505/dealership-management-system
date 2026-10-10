"""Excel import parsing + export/template builders (openpyxl, header-tolerant)."""

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _norm(header: Any) -> str:
    if header is None:
        return ""
    text = str(header).strip().lower().replace(".", " ").replace("_", " ").replace("-", " ")
    return " ".join(text.split())


def load_rows(data: bytes) -> list[dict[str, Any]]:
    """First non-empty row is the header; every following row becomes a dict keyed by normalised header."""
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    while rows and all(value is None or str(value).strip() == "" for value in rows[0]):
        rows = rows[1:]
    if not rows:
        return []
    headers = [_norm(h) for h in rows[0]]
    parsed: list[dict[str, Any]] = []
    for row in rows[1:]:
        if all(value is None or str(value).strip() == "" for value in row):
            continue
        parsed.append({headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]})
    return parsed


def pick(row: dict[str, Any], *candidates: str) -> Any:
    """First non-empty value among candidate (normalised) header names."""
    for candidate in candidates:
        value = row.get(candidate)
        if value is not None and value != "":
            return value
    return None


def as_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def as_int(value: Any, default: int = 0) -> int:
    if value is None or value == "":
        return default
    try:
        return int(float(str(value).replace(",", "").strip()))
    except (TypeError, ValueError):
        return default


def as_float(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(str(value).replace(",", "").replace("₹", "").replace("Rs.", "").strip())
    except (TypeError, ValueError):
        return default


def as_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = as_str(value)
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return text[:10]


HEADER_FILL = PatternFill("solid", fgColor="0F172A")


def build_xlsx(title: str, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        sheet.column_dimensions[get_column_letter(col)].width = widths[col - 1] if widths else 20
    for row_index, row in enumerate(rows, 2):
        for col, value in enumerate(row, 1):
            sheet.cell(row=row_index, column=col, value=value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
